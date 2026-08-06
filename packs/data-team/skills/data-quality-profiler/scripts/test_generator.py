#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Plain-assert unit tests for scripts/generate_reports.py.

Run with:
    python scripts/test_generator.py

These tests are the real regression guard for two prior real-run bugs:
  - NUL-byte padding corrupting outputs (hand-NUL-padded files slipped through).
  - "2-of-4" silent partial delivery (only some files emitted, AI shipped anyway).

No pytest framework required; uses plain `assert`. Exits 0 on success, 1 on failure.
"""
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

# Make the sibling module importable.
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import generate_reports as gr  # noqa: E402


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _synthetic_results() -> dict:
    """Minimal but schema-complete results.json: 2 tables, 3 cols each, scores,
    stats, relations, findings. Covers every field the four-file generator reads."""
    return {
        "schema_version": "1.0",
        "database": {"type": "MySQL", "name": "crm_demo"},
        "tables": [
            {
                "name": "member",
                "row_count": 120000,
                "comment": "会员表",
                "columns": [
                    {"name": "id", "type": "BIGINT", "nullable": False, "comment": "主键"},
                    {"name": "mobile", "type": "VARCHAR(20)", "nullable": True, "comment": "手机号"},
                    {"name": "status", "type": "TINYINT", "nullable": True, "comment": "状态"},
                ],
                "scores": {"完整性": 92, "唯一性": 88, "有效性": 75, "一致性": 80, "及时性": 70},
                "findings": [
                    {"severity": "yellow", "metric": "完整性", "desc": "37% 会员未绑手机", "so_what": "短信渠道对该部分不可达"},
                    {"severity": "green", "metric": "唯一性", "desc": "主键无重复", "so_what": "可作关联键"},
                ],
            },
            {
                "name": "orders",
                "row_count": 9800000,
                "comment": "订单表",
                "columns": [
                    {"name": "id", "type": "BIGINT", "nullable": False, "comment": "主键"},
                    {"name": "member_id", "type": "BIGINT", "nullable": False, "comment": "会员ID"},
                    {"name": "amount", "type": "DECIMAL(12,2)", "nullable": True, "comment": "金额"},
                ],
                "scores": {"完整性": 95, "唯一性": 99, "有效性": 90, "一致性": 60, "及时性": 85},
                "findings": [
                    {"severity": "red", "metric": "一致性", "desc": "5% 订单 member_id 孤儿", "so_what": "JOIN 前需过滤孤儿"},
                ],
            },
        ],
        "relations": [
            {"from": "orders.member_id", "to": "member.id", "confidence": 0.95, "evidence": "数据探针命中率 95%"},
        ],
        "business_stats": {
            "member_axis": [
                {"label": "会员单", "value": 7600000},
                {"label": "游客单", "value": 2200000},
            ],
            "channel_axis": [
                {"channel": "天猫", "count": 3000000, "gmv": 98000000.0},
                {"channel": "京东", "count": 2000000, "gmv": 60000000.0},
                {"channel": "自营", "count": 4800000, "gmv": 150000000.0},
            ],
            "monthly_trend": [
                {"month": "2024-07", "orders": 410000},
                {"month": "2024-08", "orders": 430000},
                {"month": "2024-09", "orders": 460000},
            ],
        },
        "overall": {
            "tables_inspected": 2,
            "red_count": 1,
            "yellow_count": 1,
            "green_count": 1,
            "summary": "整体数据质量中等，存在孤儿订单风险。",
        },
    }


def _write_results(tmpdir: Path, results: dict) -> Path:
    p = tmpdir / "results.json"
    p.write_text(json.dumps(results, ensure_ascii=False), encoding="utf-8")
    return p


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

def test_strict_schema_rejects_bad_input(tmpdir: Path) -> None:
    bad = {"tables": "not-a-list"}
    p = _write_results(tmpdir, bad)
    try:
        gr.load_and_validate(p)
    except gr.SchemaError:
        return
    raise AssertionError("expected SchemaError for malformed results.json")


def test_all_four_files_emit_and_pass_verify(tmpdir: Path) -> None:
    results = _synthetic_results()
    inp = _write_results(tmpdir, results)
    outdir = tmpdir / "reports"

    gr.main(["--input", str(inp), "--outdir", str(outdir)])

    expected = [
        "数据字典.xlsx",
        "业务统计结果.xlsx",
        "整体数据质量分析报告.docx",
        "各表体检明细.docx",
    ]
    for name in expected:
        fp = outdir / name
        assert fp.exists(), f"missing output: {name}"
        assert fp.stat().st_size > 1024, f"{name} too small ({fp.stat().st_size}B)"

    # verify_outputs must report clean.
    report = gr.verify_outputs(outdir)
    assert report["ok"], f"verify_outputs failed: {report}"
    assert set(report["files"].keys()) == set(expected)


def test_outputs_open_with_libs(tmpdir: Path) -> None:
    results = _synthetic_results()
    inp = _write_results(tmpdir, results)
    outdir = tmpdir / "reports"
    gr.main(["--input", str(inp), "--outdir", str(outdir)])

    import openpyxl
    from docx import Document

    wb1 = openpyxl.load_workbook(outdir / "数据字典.xlsx")
    assert "member" in wb1.sheetnames
    wb2 = openpyxl.load_workbook(outdir / "业务统计结果.xlsx")
    assert len(wb2.sheetnames) >= 1
    d1 = Document(str(outdir / "整体数据质量分析报告.docx"))
    assert len(d1.paragraphs) >= 1
    d2 = Document(str(outdir / "各表体检明细.docx"))
    assert len(d2.paragraphs) >= 1


def test_no_nul_bytes_in_outputs(tmpdir: Path) -> None:
    results = _synthetic_results()
    inp = _write_results(tmpdir, results)
    outdir = tmpdir / "reports"
    gr.main(["--input", str(inp), "--outdir", str(outdir)])

    for fp in outdir.iterdir():
        if not fp.is_file():
            continue
        assert not gr.has_nul_padding(fp), f"NUL padding found in {fp}"


def test_assert_no_nul_bytes_flags_corrupt_fixture(tmpdir: Path) -> None:
    """Regression for the prior real-run NUL-byte bug: a hand-NUL-padded file
    MUST be flagged by assert_no_nul_bytes."""
    bad = tmpdir / "corrupt.xlsx"
    # Real xlsx header bytes followed by NUL padding.
    bad.write_bytes(b"PK\x03\x04" + b"\x00" * 2048)
    try:
        gr.assert_no_nul_bytes(bad)
    except gr.CorruptOutputError:
        return
    raise AssertionError("assert_no_nul_bytes failed to flag NUL-padded file")


def test_re_emit_recovers_missing_file(tmpdir: Path) -> None:
    """Regression for the "2-of-4" silent partial delivery bug: if a file is
    missing/corrupt after first run, verify_outputs must re-emit it and final
    state must be all-green (or raise, never silent)."""
    results = _synthetic_results()
    inp = _write_results(tmpdir, results)
    outdir = tmpdir / "reports"
    gr.main(["--input", str(inp), "--outdir", str(outdir)])

    # Sabotage one file -> corrupt (NUL padding).
    victim = outdir / "数据字典.xlsx"
    victim.write_bytes(b"PK\x03\x04" + b"\x00" * 2048)

    # Re-run should detect + re-emit and end clean.
    gr.main(["--input", str(inp), "--outdir", str(outdir)])
    report = gr.verify_outputs(outdir)
    assert report["ok"], f"re-emit did not recover: {report}"


def test_no_leftover_chart_pngs(tmpdir: Path) -> None:
    """Regression for the e2e-flagged cosmetic issue: after main() the outdir
    must contain exactly the four named deliverables and no intermediate
    chart PNGs (prefixed with '_')."""
    results = _synthetic_results()
    inp = _write_results(tmpdir, results)
    outdir = tmpdir / "reports"
    gr.main(["--input", str(inp), "--outdir", str(outdir)])

    files = {p.name for p in outdir.iterdir() if p.is_file()}
    expected = {
        "数据字典.xlsx",
        "业务统计结果.xlsx",
        "整体数据质量分析报告.docx",
        "各表体检明细.docx",
    }
    assert files == expected, f"unexpected files in outdir: {files - expected} (expected exactly the 4 deliverables)"
    leftovers = [n for n in files if n.startswith("_")]
    assert not leftovers, f"intermediate chart PNGs leaked into outdir: {leftovers}"


def test_refuses_repo_root_outdir(tmpdir: Path) -> None:
    results = _synthetic_results()
    inp = _write_results(tmpdir, results)
    # An outdir inside a fake repo root containing .git must be refused.
    repo_root = tmpdir / "repo"
    (repo_root / ".git").mkdir(parents=True)
    outdir = repo_root / "reports"
    try:
        gr.main(["--input", str(inp), "--outdir", str(outdir)])
    except gr.UnsafePathError:
        return
    raise AssertionError("expected UnsafePathError for outdir inside repo root with .git")


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------

def _run_all() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="dqtest_"))
    try:
        tests = [
            test_strict_schema_rejects_bad_input,
            test_all_four_files_emit_and_pass_verify,
            test_outputs_open_with_libs,
            test_no_nul_bytes_in_outputs,
            test_assert_no_nul_bytes_flags_corrupt_fixture,
            test_re_emit_recovers_missing_file,
            test_no_leftover_chart_pngs,
            test_refuses_repo_root_outdir,
        ]
        passed = 0
        failed = 0
        for t in tests:
            # each test gets its own sub-tmpdir so they are independent
            sub = Path(tempfile.mkdtemp(prefix=f"{t.__name__}_"))
            try:
                t(sub)
                print(f"[PASS] {t.__name__}")
                passed += 1
            except Exception as e:  # noqa: BLE001
                print(f"[FAIL] {t.__name__}: {e}")
                failed += 1
            finally:
                shutil.rmtree(sub, ignore_errors=True)
        print(f"\n{passed} passed, {failed} failed")
        return 0 if failed == 0 else 1
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(_run_all())
